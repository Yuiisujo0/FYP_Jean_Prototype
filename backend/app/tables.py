"""Reads an uploaded CSV or Excel file into (columns, rows). Every cell becomes a string, as the engine expects."""
import base64
import binascii
import csv
import io
import re
import zipfile
from datetime import date, datetime, time

MAX_ROWS = 50_000
Row = dict[str, str]


def norm(s: str) -> str:
    """Case- and punctuation-insensitive key, matching the frontend's `norm`."""
    return re.sub(r"[^a-z0-9]+", "_", s.strip().lower()).strip("_")


def strip_ext(name: str) -> str:
    return re.sub(r"\.[A-Za-z0-9]{1,5}$", "", name)


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else format(v, ".15g")
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.isoformat(sep=" ", timespec="seconds")
    if isinstance(v, (date, time)):
        return v.isoformat()
    return str(v).strip()


def table_from_matrix(matrix: list[list[str]]) -> tuple[list[str], list[Row]]:
    """First non-empty row is the header. Blank rows are skipped. Trailing empty columns are dropped."""
    it = iter(matrix)
    header: list[str] | None = None
    for r in it:
        if any(c.strip() for c in r):
            header = [c.strip() for c in r]
            break
    if header is None:
        raise ValueError("The file has no header row.")
    while header and not header[-1]:
        header.pop()
    columns = [h or f"column_{i + 1}" for i, h in enumerate(header)]
    if len({c.lower() for c in columns}) != len(columns):
        raise ValueError("Column names must be unique.")
    rows: list[Row] = []
    for r in it:
        if not any(c.strip() for c in r):
            continue
        if len(rows) >= MAX_ROWS:
            raise ValueError(f"Too many rows (limit {MAX_ROWS}).")
        rows.append({col: (r[i].strip() if i < len(r) else "") for i, col in enumerate(columns)})
    return columns, rows


def parse_csv(text: str, delimiter: str = ",") -> tuple[list[str], list[Row]]:
    try:
        # newline="" lets the reader handle \r, \n and \r\n line endings (and newlines inside quoted cells) itself.
        matrix = [row for row in csv.reader(io.StringIO(text.lstrip("\ufeff"), newline=""), delimiter=delimiter)]
    except csv.Error as e:
        raise ValueError(f"Could not read this CSV file: {e}") from e
    return table_from_matrix(matrix)


def parse_xlsx(data: bytes) -> tuple[list[str], list[Row], str]:
    from openpyxl import load_workbook  # imported here so CSV-only use does not pay for it

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError) as e:
        raise ValueError("This is not a valid .xlsx file. Older .xls files are not supported; save it as .xlsx or CSV.") from e
    try:
        last_error = "The workbook has no data."
        for ws in wb.worksheets:
            matrix = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
            try:
                columns, rows = table_from_matrix(matrix)
            except ValueError as e:
                last_error = f"Sheet “{ws.title}”: {e}"
                continue
            return columns, rows, ws.title
        raise ValueError(last_error)
    finally:
        wb.close()


def parse_upload(filename: str, text: str, b64: str, delimiter: str = ",") -> tuple[list[str], list[Row], str | None]:
    """Returns (columns, rows, sheet name or None). Excel is recognised by the file extension."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        if not b64:
            raise ValueError("The Excel file was not received.")
        try:
            data = base64.b64decode(b64, validate=True)
        except (binascii.Error, ValueError) as e:
            raise ValueError("The Excel file was damaged in transit.") from e
        return parse_xlsx(data)
    if filename.lower().endswith(".xls"):
        raise ValueError("Older .xls files are not supported. Save it as .xlsx or CSV.")
    columns, rows = parse_csv(text, delimiter)
    return columns, rows, None
